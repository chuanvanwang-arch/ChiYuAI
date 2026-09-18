# -*- coding: utf-8 -*-
"""构建企业AI销售决策平台介绍 PPT（六大亮点版 · 12 页）
叙事主轴 = 三大业务亮点 + 三大技术亮点：
  业务：①集成大量销售成熟方法论 ②植入销售人员行为习惯 ③LTC 流程内嵌化
  技术：④上下文图谱构建客户记忆系统 ⑤销售决策模型可跟踪每个智能体任务
       ⑥复盘智能体每日自动复盘
配色与字体同现有演示 HTML（深色主题）。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------- 设计 token（与演示 HTML 同源）----------
BG    = RGBColor(0x0c,0x0e,0x14)
BG2   = RGBColor(0x13,0x16,0x1e)
CARD  = RGBColor(0x1a,0x1e,0x2a)
CYAN  = RGBColor(0x22,0xd3,0xee)
AMBER = RGBColor(0xfb,0xbf,0x24)
GREEN = RGBColor(0x34,0xd3,0x99)
RED   = RGBColor(0xf8,0x71,0x71)
PURPLE= RGBColor(0xa7,0x8b,0xfa)
BLUE  = RGBColor(0x60,0xa5,0xfa)
WHITE = RGBColor(0xe8,0xea,0xed)
GRAY  = RGBColor(0x9a,0xa0,0xb0)
DIM   = RGBColor(0x5c,0x62,0x78)
BORDER= RGBColor(0x25,0x2a,0x3a)
FONT  = 'Microsoft YaHei'

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def txt(slide, l,t,w,h, text, size=18, color=WHITE, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.2, font=FONT):
    tb = slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i,line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = font
    return tb

def card(slide, l,t,w,h, fill=CARD, line=BORDER, lw=1.0, radius=0.08):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line; sh.line.width = Pt(lw)
    try: sh.adjustments[0] = radius
    except Exception: pass
    sh.shadow.inherit = False
    return sh

def pill(slide, l,t,w,h, text, color=CYAN, fill=BG2, size=12):
    sh = card(slide,l,t,w,h,fill=fill,line=color,lw=1.0,radius=0.5)
    tf = sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=text; r.font.size=Pt(size); r.font.bold=True
    r.font.color.rgb=color; r.font.name=FONT
    return sh

def header(slide, act, active_name, title, with_line=True):
    txt(slide,0.6,0.45,3,0.4, act, 13, CYAN, True)
    txt(slide,0.6,0.82,11.5,0.95, title, 28, WHITE, True)
    txt(slide,0.6,1.72,11.5,0.4, active_name, 14, GRAY)

def newslide():
    s = prs.slides.add_slide(BLANK); bg(s); return s

BIZ = CYAN
TECH = PURPLE

# =====================================================================
# 1 · 封面
# =====================================================================
s = newslide()
txt(s,0.9,1.3,6,0.5,'AI-NATIVE SALES PLATFORM · 2026',13,CYAN,True)
txt(s,0.9,1.85,12,1.4,'企业AI销售决策平台',44,WHITE,True)
txt(s,0.9,3.3,11,0.7,'销售方法论 × 行为习惯 × LTC 内嵌 · 记忆系统 × 决策透明 × 每日复盘',17,AMBER,True)
txt(s,0.9,4.25,11,1.5,
    '不是替换原有 CRM，而是与主流 CRM 双向打通、数据回流，在其上叠加「决策智能层」：\n把成熟的销售经验、行为标准与业务流程，连同智能体的记忆、决策与复盘能力，一起装进平台。',
    15,GRAY,False,spacing=1.35)
pill(s,0.9,6.15,3.4,0.55,'业务亮点 3 + 技术亮点 3',CYAN,BG2,12)
pill(s,4.5,6.15,4.4,0.55,'15 个方法论 SKILL · 行为标准 · LTC 全链路',GREEN,BG2,12)
pill(s,9.1,6.15,3.3,0.55,'决策可溯源 · 每日自动复盘',PURPLE,BG2,12)

# =====================================================================
# 2 · 导读 AGENDA（亮点叙事）
# =====================================================================
s = newslide()
header(s,'导读 · AGENDA','本次交流，五件事','从行业之痛，到六大亮点，再到价值闭环')
agenda = [
    ('01','行业背景','B2B 销售：复杂之下，方法·习惯·流程为何流失',GRAY),
    ('02','业务亮点 ×3','成熟方法论 · 行为习惯植入 · LTC 流程内嵌',BIZ),
    ('03','技术亮点 ×3','上下文记忆 · 决策可跟踪 · 每日自动复盘',TECH),
    ('04','飞轮与价值','六大亮点如何合成能力飞轮',GREEN),
    ('05','落地路径','一键召唤 → 数据自驱 → 越用越聪明',AMBER),
]
x = 0.9
for n,t1,t2,c in agenda:
    card(s,x,2.35,2.28,3.3,fill=BG2,line=c)
    txt(s,x+0.2,2.6,2.0,0.7,n,32,c,True)
    txt(s,x+0.2,3.5,1.9,0.5,t1,17,WHITE,True)
    txt(s,x+0.2,4.15,1.95,1.3,t2,12.5,GRAY,False,spacing=1.2)
    x += 2.44

# =====================================================================
# 3 · 行业背景：方法/习惯/流程为什么必须被"平台化"
# =====================================================================
s = newslide()
header(s,'01 · 行业背景','B2B 销售三大隐性流失：方法、习惯、流程','销售复杂度之下，最值钱的三样东西最难传承')
rows = [
    ('方法论流失','报价靠拍脑袋、机会真伪凭感觉——资深销售的 BANT/MEDDICC 级判断留在人脑，人走即失。',RED),
    ('行为习惯流失','拜访无规范、跟进看心情，21 条行为标准没有强制，好习惯沉淀不下来。',AMBER),
    ('LTC 流程断裂','线索→成交全链路依赖人肉接力，任一路段断裂无人感知，商机静默流失。',RED),
]
positions = [(0.9,2.5),(6.5,2.5),(0.9,4.4)]
for (t,d,c),(l,ty) in zip(rows, positions):
    card(s,l,ty,5.4,1.7,fill=BG2,line=c)
    txt(s,l+0.25,ty+0.22,5.0,0.5,t,18,c,True)
    txt(s,l+0.25,ty+0.9,5.0,0.75,d,13.5,GRAY,False,spacing=1.25)
txt(s,0.9,6.3,11.5,0.6,'结论：销售管理要真正提效，不是把 CRM 表单搬上线，而是把「方法、习惯、流程」变成平台的系统能力。',15,CYAN,True)

# =====================================================================
# 4 · 业务亮点一：集成大量销售成熟方法论
# =====================================================================
s = newslide()
header(s,'02 · 业务亮点 ①','集成大量销售成熟方法论：判断不再拍脑袋','15 个方法论独立 SKILL，覆盖线索到成交的每一步')
methods = [
    ('BANT','预算/权限/需求/时间线','线索初筛',BIZ),
    ('MEDDICC','痛点/经济买家/决策流程','机会真假',PURPLE),
    ('机会矩阵','价值 × 赢率','资源投入',BLUE),
    ('角色地图','决策链/影响者/使用者','主攻谁',AMBER),
    ('三级报价','开盘/目标/底价+条件换折扣','报价底线',GREEN),
    ('风险权衡','风险 × 收益/红线','签不签单',RED),
    ('止损点','投入预算/退出门','何时退出',AMBER),
    ('事实vs话术','事实/证据/话术/异议','不被忽悠',GREEN),
]
positions = [(0.9,2.35),(4.35,2.35),(7.8,2.35),(11.25,2.35),
             (0.9,3.85),(4.35,3.85),(7.8,3.85),(11.25,3.85)]
for (t,d,u,c),(l,ty) in zip(methods, positions):
    card(s,l,ty,3.2,1.4,fill=CARD,line=c)
    txt(s,l+0.18,ty+0.13,2.85,0.5,t,16,c,True)
    txt(s,l+0.18,ty+0.65,2.85,0.45,d,10.5,GRAY,False,spacing=1.05)
    txt(s,l+0.18,ty+1.02,2.85,0.35,f'用途：{u}',10,DIM,False,spacing=1.0)
txt(s,0.9,5.5,11.5,1.5,
    '更多：method-funnel-classification（漏斗分类）· method-stage-progression（阶段推进）· method-behavior-standard（行为标准）· method-intake-routing（接诊分流）· method-quote-engine（报价引擎）· method-review-gate（评审闸门）…\n\n每个决策「问什么 / 要什么输入 / 用什么方法 / 结论留痕」，方法论不再是墙上的字，而是每一次判断的默认动作。',
    13,GRAY,False,spacing=1.3)
pill(s,0.9,6.85,5.0,0.5,'BANT/MEDDICC/机会矩阵/角色地图/三级报价 已落地',GREEN,BG2,12)

# =====================================================================
# 5 · 业务亮点二：植入销售人员行为习惯
# =====================================================================
s = newslide()
header(s,'02 · 业务亮点 ②','植入销售人员行为习惯：好销售不是天赋，是可执行的标准','销售管理体系行为标准 SKILL 化，21 条可测、可查、可复盘')
habits = [
    ('21 条行为标准','按销售任务分类，每条都可测评、可打分',GREEN),
    ('拜访活动管理','拜访前准备 → 拜访中记录 → 拜访后复盘，闭环留痕',BIZ),
    ('商机阶段动作','每一阶段该做什么、做到什么算达标，有标尺',BLUE),
    ('行为达标看板','「我的行为」实时可见，销售自驱改进',AMBER),
]
positions = [(0.9,2.4),(6.5,2.4),(0.9,4.15),(6.5,4.15)]
for (t,d,c),(l,ty) in zip(habits, positions):
    card(s,l,ty,5.4,1.6,fill=BG2,line=c)
    txt(s,l+0.25,ty+0.2,5.0,0.5,t,17,c,True)
    txt(s,l+0.25,ty+0.85,5.0,0.6,d,13,GRAY,False,spacing=1.2)
txt(s,0.9,5.95,11.5,0.9,
    '系统实时对照行为标准：拜访频次够不够、商机阶段停留是否超阈值、合格线是否达成。\n习惯不再是口号，而是逐条可测量、可检查、可复盘的执行标准。',
    14,GRAY,False,spacing=1.3)
pill(s,0.9,6.9,7.0,0.5,'行为标准 · 21 条合格线 · 拜访闭环 · 达标看板',AMBER,BG2,12)

# =====================================================================
# 6 · 业务亮点三：LTC 流程内嵌化
# =====================================================================
s = newslide()
header(s,'02 · 业务亮点 ③','LTC 流程内嵌化：全链路在一条轨道上跑，不再人肉接力','线索 → 商机 → 方案 → 报价 → 合同 → 交付 → 回款 → 复盘')
ltc = [
    ('线索',   '培育 · 初筛 · 升级',  GRAY),
    ('商机',   'MEDDICC · 分级 · 阶段推进', BIZ),
    ('方案',   '需求补全 · 分层 · 对齐', BLUE),
    ('报价',   '三级报价 · 成本透视', PURPLE),
    ('合同',   '评审闸门 · 风险权衡', RED),
    ('交付',   '实施排期 · 链断裂预警', AMBER),
    ('回款',   '逾期分级 · 差异化催收', GREEN),
    ('复盘',   '订单复盘 · 复购转化', AMBER),
]
x=0.9
for t,d,c in ltc:
    card(s,x,2.7,1.42,1.5,fill=BG2,line=c)
    txt(s,x+0.12,2.85,1.2,0.45,t,13.5,c,True,align=PP_ALIGN.CENTER)
    txt(s,x+0.12,3.4,1.2,0.75,d,9.5,GRAY,False,align=PP_ALIGN.CENTER,spacing=1.05)
    if t!='复盘':
        txt(s,x+1.42,2.95,0.14,0.6,'→',14,DIM,True,align=PP_ALIGN.CENTER)
    x += 1.56
txt(s,0.9,4.5,11.5,1.6,
    '流程内嵌 ≠ 表单录账：每个环节的方法论自动附着、每个节点的判断自动留痕、每段链路断裂自动预警。\n\n· 商机 30 天无方案 → 预警\n· 回款逾期 → 分级催收\n· 阶段只进不退 → 异常探测',
    14,GRAY,False,spacing=1.3)
pill(s,0.9,6.3,4.2,0.5,'L2C 管道全流程受控',BIZ,BG2,12)
pill(s,5.3,6.3,5.4,0.5,'链路断裂主动探测（SSE 推送）',RED,BG2,12)

# =====================================================================
# 7 · 技术亮点一：上下文图谱构建客户记忆系统
# =====================================================================
s = newslide()
header(s,'03 · 技术亮点 ①','上下文图谱构建客户记忆系统：从「数据库」到「记忆系统」','知识图谱 + 向量 + 决策事件，客户全景可追溯')
card(s,0.9,2.45,5.4,2.9,fill=BG2,line=TECH)
txt(s,1.15,2.65,5.0,0.5,'上下文图谱（KG + 向量）',17,TECH,True)
txt(s,1.15,3.3,5.0,1.8,
    '· 写库即构建：数据写入的同时生成本体关系与向量索引\n· 客户/商机/报价/合同/回款…不再是孤立表单，而是连成知识网\n· 语义检索：「上次那套系统再来一套」→ 秒级命中历史',
    13.5,WHITE,False,spacing=1.3)
card(s,6.5,2.45,5.4,2.9,fill=BG2,line=BIZ)
txt(s,6.75,2.65,5.0,0.5,'五类高频记忆自动沉淀',17,BIZ,True)
txt(s,6.75,3.3,5.0,1.8,
    '① 客户业务（行业/KPI/痛点/预算）\n② 组织决策链（架构/角色/内线）\n③ 竞争情报（对手方案/报价）\n④ 我方内部（成本/毛利红线）\n⑤ 风险（经营/合同/交付/回款）',
    13.5,WHITE,False,spacing=1.25)
txt(s,0.9,5.6,11.5,1.2,
    '判据：只答 What 是数据库；能答 Why（谁/何时/结论/理由）才是记忆系统。\n客户 360 正向追溯 + 反向溯源：一笔报价能追到它的成本结构与同类先例。',
    14,AMBER,False,spacing=1.3)

# =====================================================================
# 8 · 技术亮点二：销售决策模型可跟踪每个智能体任务
# =====================================================================
s = newslide()
header(s,'03 · 技术亮点 ②','销售决策模型可跟踪每个智能体任务：每一步都有据可查','决策全链路留痕：智能体 → 任务 → 决策 → 结果 → 回流')
card(s,0.9,2.45,5.4,2.9,fill=BG2,line=TECH)
txt(s,1.15,2.65,5.0,0.5,'决策事件为主轴',17,TECH,True)
txt(s,1.15,3.3,5.0,1.8,
    '· 每个销售判断 = 一条决策记录（谁/何时/结论/理由）\n· 决策边七类：针对/参考先例/由异常触发/确立标杆/推翻/直接引发/间接影响\n· 写操作必经「决策第 0 闸」：先决策、后写入，全程携带决策 ID',
    13.5,WHITE,False,spacing=1.3)
card(s,6.5,2.45,5.4,2.9,fill=BG2,line=BIZ)
txt(s,6.75,2.65,5.0,0.5,'智能体任务可跟踪',17,BIZ,True)
txt(s,6.75,3.3,5.0,1.8,
    '· 接诊/报价/跟进/评审各有专职智能体，谁在干活一目了然\n· 每个任务绑定商机、决策与执行记录，卡在哪条链路可追溯\n· 商机分级：一般全自动跑 / 普通人机协同 / 重大两级评审',
    13.5,WHITE,False,spacing=1.3)
txt(s,0.9,5.6,11.5,1.2,
    '结果回流校准：已执行决策的输赢/回款结果回流系统，自动识别「偏乐观判断」并下调置信度——错的判断被识别，对的判断被强化。',
    14,AMBER,False,spacing=1.3)

# =====================================================================
# 9 · 技术亮点三：复盘智能体每日自动复盘
# =====================================================================
s = newslide()
header(s,'03 · 技术亮点 ③','复盘智能体每日自动复盘：经验每天被沉淀，不等月底','夜间批量扫描 + LLM 深度归因 + 可审查的处方草稿')
card(s,0.9,2.45,5.4,2.9,fill=BG2,line=TECH)
txt(s,1.15,2.65,5.0,0.5,'每日自动复盘（retro 引擎）',17,TECH,True)
txt(s,1.15,3.3,5.0,1.8,
    '· 每日全量扫描当日决策集群，自动归因沉淀\n· 七类根因分类：字段错配/信息不全/输入不及时/维度缺失/边缺失/次序不当/先例污染\n· 输出可审查的复盘报告（draft_patches），不经审批不自动改',
    13.5,WHITE,False,spacing=1.3)
card(s,6.5,2.45,5.4,2.9,fill=BG2,line=BIZ)
txt(s,6.75,2.65,5.0,0.5,'复盘闭环（飞轮）',17,BIZ,True)
txt(s,6.75,3.3,5.0,1.8,
    '· 丢单复盘 → 教训变成 SOP（订单复盘案例）\n· 成交复盘 → 同类项目报价锚点 + 复购转化排程\n· 复盘报告进入决策记忆，下次同类判断自动调用',
    13.5,WHITE,False,spacing=1.3)
txt(s,0.9,5.6,11.5,1.2,
    '案例实证：成交 ¥12.6 万年度系统单，复盘出「需求变更未设闸门 + 验收未锁死」两类根因 → 沉淀为变更收费/验收即锁 SOP → 客户满意转复购。',
    14,GREEN,False,spacing=1.3)

# =====================================================================
# 10 · 全景：六大亮点合成飞轮
# =====================================================================
s = newslide()
header(s,'04 · 全景','六大亮点不是六个功能，而是一个能力飞轮','方法论 → 行为 → 流程 → 记忆 → 决策 → 复盘，环环咬合')
fly = [
    ('业务 · 方法论','15 个 method-* SKILL：判断有方法',BIZ),
    ('业务 · 行为习惯','21 条行为标准：动作有标尺',AMBER),
    ('业务 · LTC 内嵌','线索到回款一条轨道：流程不脱节',GREEN),
    ('技术 · 记忆系统','上下文图谱：客户全景不丢失',TECH),
    ('技术 · 决策可跟踪','决策第 0 闸 + 留痕：判断可溯源',BLUE),
    ('技术 · 每日复盘','retro 引擎：经验每天被沉淀',GREEN),
]
positions = [(0.9,2.4),(5.1,2.4),(9.3,2.4),(0.9,4.15),(5.1,4.15),(9.3,4.15)]
for (t,d,c),(l,ty) in zip(fly, positions):
    card(s,l,ty,3.8,1.6,fill=BG2,line=c)
    txt(s,l+0.25,ty+0.2,3.3,0.5,t,15.5,c,True)
    txt(s,l+0.25,ty+0.85,3.3,0.65,d,12.5,GRAY,False,spacing=1.15)
txt(s,0.9,5.95,11.5,1.0,
    '越用越聪明：捕获 Capture → 拼装 Assemble → 决策 Decide → 回流 Feedback → 复盘 Review。\n每一圈，先例更丰富、判断更准、人更省心。',
    15,WHITE,False,spacing=1.35)

# =====================================================================
# 11 · 价值与落地
# =====================================================================
s = newslide()
header(s,'04 · 价值与落地','用真实业务印证六点落地','报价 / 需求补全 / 链预警 / 复盘 / 角色自适应，全部来自真实业务文档')
proofs = [
    ('系统方案报价','一句话报出 A/B 两方案，成本结构、毛利底线全透视',GREEN),
    ('需求自动补全','8 项必填缺失自动追问，一个微信来回齐',BIZ),
    ('交付风险预警','链断裂/回款逾期先于提问主动推送',RED),
    ('决策全程留痕','谁/何时/结论/理由可溯源，可复盘',TECH),
    ('角色自适应','同一句话，销售/经理/高管/财务各见所需',AMBER),
    ('订单复盘复购','¥12.6 万复盘成 SOP，复购自动排程',GREEN),
]
positions = [(0.9,2.4),(6.5,2.4),(0.9,4.15),(6.5,4.15),(0.9,5.9),(6.5,5.9)]
for (t,d,c),(l,ty) in zip(proofs, positions):
    card(s,l,ty,5.4,1.55,fill=CARD,line=c)
    txt(s,l+0.25,ty+0.16,5.0,0.45,t,15,c,True)
    txt(s,l+0.25,ty+0.7,5.0,0.75,d,12.5,GRAY,False,spacing=1.15)

# =====================================================================
# 12 · 结语
# =====================================================================
s = newslide()
header(s,'结语 · WHAT\'S NEXT','让「方法、习惯、流程、记忆、决策、复盘」都成为系统能力','建议下一步：用真实数据说话')
txt(s,0.9,2.5,11.5,1.2,
    '平台不是工具，而是会自检、会建议、会进化的销售中枢。\n先跑通「一般商机全自动闭环」：自动接诊、标准报价、自动跟进，让销售立刻减负。',
    17,WHITE,False,spacing=1.35)
pill(s,0.9,4.1,3.5,0.6,'第一步 · 一键召唤专家包',CYAN,BG2,13)
pill(s,4.7,4.1,3.8,0.6,'第二步 · 方法论+行为标准铺开',AMBER,BG2,13)
pill(s,8.6,4.1,3.8,0.6,'第三步 · 每日复盘自动运转',GREEN,BG2,13)
txt(s,0.9,5.35,11.5,1.3,
    '企业AI销售决策平台 —— 不替换原有 CRM，与主流 CRM 双向打通、叠加决策智能层；\n新增经销商管理、销售拓客、首次登录打通（邮箱/日历/会议/微信）、日期驱动自动执行。\n让每一次报价有方法、每一次拜访有标准、每一个决策可追溯、每一天的经验都被沉淀。',
    16,AMBER,True,spacing=1.3)

# ---------- 保存 ----------
out = 'D:/system/CRM-ai-native/doc/企业AI销售决策平台介绍_2026.pptx'
try:
    prs.save(out)
    print('SAVED_OK', out, 'slides=', len(prs.slides._sldIdLst))
except PermissionError:
    out2 = 'D:/system/CRM-ai-native/doc/企业AI销售决策平台介绍_2026_tmp.pptx'
    prs.save(out2)
    print('LOCKED_SAVED_AS', out2, 'slides=', len(prs.slides._sldIdLst))