# -*- coding: utf-8 -*-
"""构建企业AI销售决策平台介绍 PPT（能力全景版 · 14 页）
叙事主轴 = 三大业务亮点 + 三大技术亮点 + 两大平台亮点：
  业务：①集成大量销售成熟方法论 ②植入销售人员行为习惯 ③LTC 流程内嵌化
  技术：④上下文图谱客户记忆 + 租户 Know-How ⑤认知决策引擎（九尺子 + 全程留痕）
       ⑥复盘智能体每日自动复盘 + 知识回写
  平台：⑦认知驱动决策引擎 K-M-D ⑧多行业零代码配置上线
配色与字体同现有演示 HTML（深色主题）。
输出：企业AI销售决策平台介绍_2026_v2.pptx（保留 2026 原版不覆盖）
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------- 设计 token（与演示 HTML 同源）----------
BG    = RGBColor(0x0c,0x0e,0x14)
BG2   = RGBColor(0x13,0x16,0x1e)
CARD  = RGBColor(0x1a,0x1e,0x2a)
CYAN  = RGBColor(0x22,0xd3,0xee)
AMBER = RGBColor(0xfb,0xbf,0x24)
GREEN = RGBColor(0x34,0xd3,0x99)
RED   = RGBColor(0xf8,0x71,0x71)
PURPLE= RGBColor(0xa7,0x8b,0xfa)
BLUE  = RGBColor(0x60,0xa5,0xfa)
WHITE = RGBColor(0xe8,0xea,0xed)
GRAY  = RGBColor(0x9a,0xa0,0xb0)
DIM   = RGBColor(0x5c,0x62,0x78)
BORDER= RGBColor(0x25,0x2a,0x3a)
FONT  = 'Microsoft YaHei'

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def txt(slide, l,t,w,h, text, size=18, color=WHITE, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.2, font=FONT):
    tb = slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i,line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = font
    return tb

def card(slide, l,t,w,h, fill=CARD, line=BORDER, lw=1.0, radius=0.08):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(l),Inches(t),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line; sh.line.width = Pt(lw)
    try: sh.adjustments[0] = radius
    except Exception: pass
    sh.shadow.inherit = False
    return sh

def pill(slide, l,t,w,h, text, color=CYAN, fill=BG2, size=12):
    sh = card(slide,l,t,w,h,fill=fill,line=color,lw=1.0,radius=0.5)
    tf = sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=text; r.font.size=Pt(size); r.font.bold=True
    r.font.color.rgb=color; r.font.name=FONT
    return sh

def header(slide, act, active_name, title, with_line=True):
    txt(slide,0.6,0.45,3,0.4, act, 13, CYAN, True)
    txt(slide,0.6,0.82,11.5,0.95, title, 28, WHITE, True)
    txt(slide,0.6,1.72,11.5,0.4, active_name, 14, GRAY)

_page = {'n': 0}
def newslide():
    s = prs.slides.add_slide(BLANK); bg(s)
    _page['n'] += 1
    txt(s,12.3,6.95,0.75,0.3, f'{_page["n"]:02d}', 10, DIM, False, align=PP_ALIGN.RIGHT)
    return s

BIZ = CYAN
TECH = PURPLE

# =====================================================================
# 1 · 封面
# =====================================================================
s = newslide()
txt(s,0.9,1.25,6,0.5,'AI-NATIVE SALES PLATFORM · 2026-09 版',13,CYAN,True)
txt(s,0.9,1.8,12,1.4,'企业AI销售决策平台',44,WHITE,True)
txt(s,0.9,3.25,11.5,0.7,
    '方法论 × 行为习惯 × LTC 内嵌 · 记忆 × 认知决策 × 每日复盘 × 多行业配置',17,AMBER,True)
txt(s,0.9,4.2,11,1.3,
    '不是把流程搬到线上，而是把成熟的销售经验、行为标准与业务流程，\n'
    '连同智能体的记忆、九尺子校验与复盘能力，一起装进平台；\n'
    '换一个行业，只换一份配置画像，不改一行代码。',
    15,GRAY,False,spacing=1.3)
pill(s,0.9,6.05,3.2,0.55,'业务亮点 3',CYAN,BG2,12)
pill(s,4.3,6.05,3.2,0.55,'技术亮点 3',PURPLE,BG2,12)
pill(s,7.7,6.05,4.6,0.55,'平台亮点 2 · K-M-D + 多行业配置',GREEN,BG2,12)

# =====================================================================
# 2 · 导读 AGENDA
# =====================================================================
s = newslide()
header(s,'导读 · AGENDA','本次交流，六件事','从行业之痛，到八大能力，再到价值闭环')
agenda = [
    ('01','行业背景','B2B 销售：方法·习惯·流程·知识为何持续流失',GRAY),
    ('02','业务亮点 ×3','成熟方法论 · 行为习惯植入 · LTC 流程内嵌',BIZ),
    ('03','技术亮点 ×3','上下文记忆 · 认知决策引擎 · 每日复盘回写',TECH),
    ('04','平台亮点 ①','K-M-D 三层管道：判断可解释、可复现',CYAN),
    ('05','平台亮点 ②','多行业零代码配置：换行业不改代码',GREEN),
    ('06','飞轮与落地','八大能力合成飞轮 · 真实业务印证',AMBER),
]
x = 0.5
for n,t1,t2,c in agenda:
    card(s,x,2.4,1.95,3.2,fill=BG2,line=c)
    txt(s,x+0.18,2.62,1.7,0.6,n,26,c,True)
    txt(s,x+0.18,3.42,1.62,0.5,t1,14,WHITE,True)
    txt(s,x+0.18,4.05,1.62,1.4,t2,11.5,GRAY,False,spacing=1.2)
    x += 2.08
txt(s,0.5,6.0,12.4,0.5,'一条主线：把「人脑里的判断力」变成「系统里的可解释能力」，再把这套能力复制到任意行业。',13,CYAN,True)

# =====================================================================
# 3 · 行业背景：四重流失
# =====================================================================
s = newslide()
header(s,'01 · 行业背景','B2B 销售四重隐性流失：方法、习惯、流程、知识','销售复杂度之下，最值钱的四样东西最难传承')
rows = [
    ('方法论流失','报价靠拍脑袋、机会真伪凭感觉——资深销售的 BANT/MEDDICC 级判断留在人脑，人走即失。',RED),
    ('行为习惯流失','拜访无规范、跟进看心情，21 条行为标准没有强制，好习惯沉淀不下来。',AMBER),
    ('LTC 流程断裂','线索→成交全链路依赖人肉接力，任一路段断裂无人感知，商机静默流失。',RED),
    ('知识无法复用','行业 Know-How 与丢单教训散落在聊天记录里，新人接手要重新踩一遍坑。',AMBER),
]
positions = [(0.9,2.45),(6.5,2.45),(0.9,4.25),(6.5,4.25)]
for (t,d,c),(l,ty) in zip(rows, positions):
    card(s,l,ty,5.4,1.65,fill=BG2,line=c)
    txt(s,l+0.25,ty+0.2,5.0,0.5,t,17,c,True)
    txt(s,l+0.25,ty+0.85,5.0,0.7,d,12.5,GRAY,False,spacing=1.25)
txt(s,0.9,6.25,11.5,0.6,'结论：销售管理要真正提效，不是把 CRM 表单搬上线，而是把「方法、习惯、流程、知识」变成平台的系统能力。',15,CYAN,True)

# =====================================================================
# 4 · 业务亮点一：集成大量销售成熟方法论
# =====================================================================
s = newslide()
header(s,'02 · 业务亮点 ①','集成大量销售成熟方法论：判断不再拍脑袋','15 个方法论独立 SKILL，覆盖线索到成交的每一步')
methods = [
    ('BANT','预算/权限/需求/时间线','线索初筛',BIZ),
    ('MEDDICC','痛点/经济买家/决策流程','机会真假',PURPLE),
    ('机会矩阵','价值 × 赢率','资源投入',BLUE),
    ('角色地图','决策链/影响者/使用者','主攻谁',AMBER),
    ('三级报价','开盘/目标/底价+条件换折扣','报价底线',GREEN),
    ('风险权衡','风险 × 收益/红线','签不签单',RED),
    ('止损点','投入预算/退出门','何时退出',AMBER),
    ('事实vs话术','事实/证据/话术/异议','不被忽悠',GREEN),
]
positions = [(0.9,2.35),(4.35,2.35),(7.8,2.35),(11.25,2.35),
             (0.9,3.85),(4.35,3.85),(7.8,3.85),(11.25,3.85)]
for (t,d,u,c),(l,ty) in zip(methods, positions):
    card(s,l,ty,3.2,1.4,fill=CARD,line=c)
    txt(s,l+0.18,ty+0.13,2.85,0.5,t,16,c,True)
    txt(s,l+0.18,ty+0.65,2.85,0.45,d,10.5,GRAY,False,spacing=1.05)
    txt(s,l+0.18,ty+1.02,2.85,0.35,f'用途：{u}',10,DIM,False,spacing=1.0)
txt(s,0.9,5.5,11.5,1.4,
    '更多：method-funnel-classification（漏斗分类）· method-stage-progression（阶段推进）· method-behavior-standard（行为标准）· method-intake-routing（接诊分流）· method-quote-engine（报价引擎）· method-review-gate（评审闸门）…\n\n每个决策「问什么 / 要什么输入 / 用什么方法 / 结论留痕」，每条方法论带权重与必填维度，缺项直接报缺。',
    12.5,GRAY,False,spacing=1.3)
pill(s,0.9,6.9,6.2,0.5,'BANT/MEDDICC/机会矩阵/角色地图/三级报价 已落地',GREEN,BG2,12)

# =====================================================================
# 5 · 业务亮点二：植入销售人员行为习惯
# =====================================================================
s = newslide()
header(s,'02 · 业务亮点 ②','植入销售人员行为习惯：好销售不是天赋，是可执行的标准','销售管理体系行为标准 SKILL 化，21 条可测、可查、可复盘')
habits = [
    ('21 条行为标准','按销售任务分类，每条都可测评、可打分',GREEN),
    ('拜访活动管理','拜访前准备 → 拜访中记录 → 拜访后复盘，闭环留痕',BIZ),
    ('商机阶段动作','每一阶段该做什么、做到什么算达标，有标尺',BLUE),
    ('行为达标看板','「我的行为」实时可见，销售自驱改进',AMBER),
]
positions = [(0.9,2.4),(6.5,2.4),(0.9,4.15),(6.5,4.15)]
for (t,d,c),(l,ty) in zip(habits, positions):
    card(s,l,ty,5.4,1.6,fill=BG2,line=c)
    txt(s,l+0.25,ty+0.2,5.0,0.5,t,17,c,True)
    txt(s,l+0.25,ty+0.85,5.0,0.6,d,13,GRAY,False,spacing=1.2)
txt(s,0.9,5.95,11.5,0.9,
    '系统实时对照行为标准：拜访频次够不够、商机阶段停留是否超阈值、合格线是否达成。\n习惯不再是口号，而是逐条可测量、可检查、可复盘的执行标准。',
    14,GRAY,False,spacing=1.3)
pill(s,0.9,6.9,7.0,0.5,'行为标准 · 21 条合格线 · 拜访闭环 · 达标看板',AMBER,BG2,12)

# =====================================================================
# 6 · 业务亮点三：LTC 流程内嵌化
# =====================================================================
s = newslide()
header(s,'02 · 业务亮点 ③','LTC 流程内嵌化：全链路在一条轨道上跑，不再人肉接力','线索 → 商机 → 方案 → 报价 → 合同 → 交付 → 回款 → 复盘')
ltc = [
    ('线索',   '培育 · 初筛 · 升级',  GRAY),
    ('商机',   'MEDDICC · 分级 · 阶段推进', BIZ),
    ('方案',   '需求补全 · 分层 · 对齐', BLUE),
    ('报价',   '三级报价 · 成本透视', PURPLE),
    ('合同',   '评审闸门 · 风险权衡', RED),
    ('交付',   '实施排期 · 链断裂预警', AMBER),
    ('回款',   '逾期分级 · 差异化催收', GREEN),
    ('复盘',   '订单复盘 · 复购转化', AMBER),
]
x=0.9
for t,d,c in ltc:
    card(s,x,2.7,1.42,1.5,fill=BG2,line=c)
    txt(s,x+0.12,2.85,1.2,0.45,t,13.5,c,True,align=PP_ALIGN.CENTER)
    txt(s,x+0.12,3.4,1.2,0.75,d,9.5,GRAY,False,align=PP_ALIGN.CENTER,spacing=1.05)
    if t!='复盘':
        txt(s,x+1.42,2.95,0.14,0.6,'→',14,DIM,True,align=PP_ALIGN.CENTER)
    x += 1.56
txt(s,0.9,4.5,11.5,1.6,
    '流程内嵌 ≠ 表单录账：每个环节的方法论自动附着、每个节点的判断自动留痕、每段链路断裂自动预警。\n\n· 商机 30 天无方案 → 预警　· 回款逾期 → 分级催收　· 阶段只进不退 → 异常探测',
    14,GRAY,False,spacing=1.3)
pill(s,0.9,6.3,4.2,0.5,'LTC 管道全流程受控',BIZ,BG2,12)
pill(s,5.3,6.3,5.4,0.5,'链路断裂主动探测（SSE 推送）',RED,BG2,12)

# =====================================================================
# 7 · 技术亮点一：上下文记忆 + 租户 Know-How
# =====================================================================
s = newslide()
header(s,'03 · 技术亮点 ①','上下文图谱 + 租户级 Know-How：从「数据库」到「记忆系统」','写库即构建，客户全景可追溯，行业 Know-How 按租户独立沉淀')
card(s,0.9,2.45,3.5,2.9,fill=BG2,line=TECH)
txt(s,1.12,2.65,3.1,0.5,'上下文图谱（KG + 向量）',15,TECH,True)
txt(s,1.12,3.25,3.1,1.9,
    '· 写库即构建：写入同时生成本体关系与向量索引\n· 客户/商机/报价/合同/回款连成知识网\n· 语义检索：「上次那套系统再来一套」秒级命中',
    12,WHITE,False,spacing=1.3)
card(s,4.7,2.45,3.5,2.9,fill=BG2,line=BIZ)
txt(s,4.92,2.65,3.1,0.5,'五类高频记忆',15,BIZ,True)
txt(s,4.92,3.25,3.1,1.9,
    '① 客户业务（行业/KPI/痛点/预算）\n② 组织决策链（架构/角色/内线）\n③ 竞争情报（对手方案/报价）\n④ 我方内部（成本/毛利红线）\n⑤ 风险（经营/合同/交付/回款）',
    12,WHITE,False,spacing=1.25)
card(s,8.5,2.45,3.5,2.9,fill=BG2,line=AMBER)
txt(s,8.72,2.65,3.1,0.5,'租户级 Know-How',15,AMBER,True)
txt(s,8.72,3.25,3.1,1.9,
    '① ICP 画像　② 竞品对比\n③ 客户异议　④ 买家语言\n\n经「决策第 0 闸 + 人工确认」写入，按租户隔离，按需注入智能体上下文。',
    12,WHITE,False,spacing=1.25)
txt(s,0.9,5.6,11.5,1.2,
    '判据：只答 What 是数据库；能答 Why（谁/何时/结论/理由）才是记忆系统。\n客户 360 正向追溯 + 反向溯源：一笔报价能追到它的成本结构与同类先例；决策链与先例可检索、可比对，新人 10 秒接手。',
    14,AMBER,False,spacing=1.3)

# =====================================================================
# 8 · 技术亮点二：认知决策引擎（九尺子）
# =====================================================================
s = newslide()
header(s,'03 · 技术亮点 ②','认知决策引擎：每一步判断都过九把尺子','8 项确定性判定（可复现）+ 1 项（清晰性）可选 LLM 精评，默认关闭')
rulers = [
    ('清晰性','目标/问题/假设措辞是否明确（无模糊词）','LLM 默认关',PURPLE),
    ('准确性','假设是否有 grounded/contradicted/unsupported 证据','确定性',BLUE),
    ('精确性','条件值是否落到具体类型、单位、时点','确定性',AMBER),
    ('相关性','供给项与必填维度（required_dims）命中率','确定性',GREEN),
    ('深度','推论链层数 ≥2 且触及根因','确定性',CYAN),
    ('广度','立场多样性（≥3 且含反方）','确定性',RED),
    ('逻辑性','每条推论同时有证据与假设，依赖图无环','确定性',AMBER),
    ('重要性','概念引用按权重加权命中率','确定性',GREEN),
    ('公平性','反面先例被检索到且被消费','确定性',BIZ),
]
positions = [(0.9,2.35),(5.1,2.35),(9.3,2.35),
             (0.9,3.72),(5.1,3.72),(9.3,3.72),
             (0.9,5.09),(5.1,5.09),(9.3,5.09)]
for (t,d,k,c),(l,ty) in zip(rulers, positions):
    card(s,l,ty,3.9,1.28,fill=CARD,line=c)
    txt(s,l+0.18,ty+0.12,2.4,0.4,t,15,c,True)
    txt(s,l+0.18,ty+0.58,3.55,0.55,d,10.5,GRAY,False,spacing=1.1)
    kind_color = PURPLE if k.startswith('LLM') else DIM
    txt(s,l+2.55,ty+0.14,1.2,0.35,k,10,kind_color,True,align=PP_ALIGN.RIGHT)
txt(s,0.9,6.55,11.5,0.55,
    '0–4 分制，及格线与权重走配置（禁硬编码）；示例：一次机会评估得「精确性 1 ⚠ / 广度 1 ⚠」→ 归因（预算区间未落到具体值、立场缺反方）→ 补齐后复核，不放行、不静默通过。',
    12.5,AMBER,True)

# =====================================================================
# 9 · 平台亮点一：K-M-D 三层管道
# =====================================================================
s = newslide()
header(s,'04 · 平台亮点 ①','认知驱动决策引擎 K-M-D：知识 → 方法 → 校验','把「拍脑袋」变成可解释、可复现、可审计的判断')
card(s,0.9,2.5,3.7,3.0,fill=BG2,line=CYAN)
txt(s,1.12,2.72,3.3,0.5,'K · 知识层 Knowledge',16,CYAN,True)
txt(s,1.12,3.35,3.3,1.9,
    '· 五类客户记忆自动装配\n· 租户 Know-How 按需注入\n· 同类先例与历史决策链\n· 缺什么信息就报什么，不做无米之炊',
    12,WHITE,False,spacing=1.3)
txt(s,4.75,3.6,0.4,0.5,'→',20,DIM,True,align=PP_ALIGN.CENTER)
card(s,5.3,2.5,3.7,3.0,fill=BG2,line=PURPLE)
txt(s,5.52,2.72,3.3,0.5,'M · 方法层 Method',16,PURPLE,True)
txt(s,5.52,3.35,3.3,1.9,
    '· 按场景装载方法论 SKILL\n· 装载行为标准与必填维度\n· 每条方法论带权重，参与判定\n· 不是把方法名写进提示词就算用上',
    12,WHITE,False,spacing=1.3)
txt(s,9.15,3.6,0.4,0.5,'→',20,DIM,True,align=PP_ALIGN.CENTER)
card(s,9.7,2.5,3.0,3.0,fill=BG2,line=AMBER)
txt(s,9.9,2.72,2.6,0.5,'D · 决策层 Decision',16,AMBER,True)
txt(s,9.9,3.35,2.6,1.9,
    '· 九尺子逐项打分（0–4）\n· 8 项确定性 + 1 项 LLM（默认关）\n· 输出结论/风险清单/止损条件\n· 携带 concept_refs 概念引用',
    12,WHITE,False,spacing=1.3)
txt(s,0.9,5.75,11.8,1.1,
    '输出即证据：intent（目的+问题）· assumptions · inference · viewpoints · implications · risk_register · stop_loss · concept_refs · rubric —— 审计时可从结论一路追回假设的证据与概念引用。\n业务事件自动触发智能体派发（decision-enrich / decision-execute），写操作必经「决策第 0 闸」，决策链与先例可检索复用。',
    13.5,GRAY,False,spacing=1.3)

# =====================================================================
# 10 · 技术亮点三：每日复盘 + 知识回写
# =====================================================================
s = newslide()
header(s,'03 · 技术亮点 ③','复盘智能体每日自动复盘：经验每天被沉淀，知识自动回写','夜间批量扫描 + LLM 深度归因 + 可审查的处方草稿')
card(s,0.9,2.45,5.4,2.9,fill=BG2,line=TECH)
txt(s,1.15,2.65,5.0,0.5,'每日自动复盘（retro 引擎）',17,TECH,True)
txt(s,1.15,3.3,5.0,1.8,
    '· 每日全量扫描当日决策集群，自动归因沉淀\n· 七类根因分类：字段错配/信息不全/输入不及时/维度缺失/边缺失/次序不当/先例污染\n· 输出可审查的复盘报告，不经审批不自动改',
    13.5,WHITE,False,spacing=1.3)
card(s,6.5,2.45,5.4,2.9,fill=BG2,line=BIZ)
txt(s,6.75,2.65,5.0,0.5,'复盘闭环（飞轮 + 知识回写）',17,BIZ,True)
txt(s,6.75,3.3,5.0,1.8,
    '· 丢单复盘 → 教训变成 SOP\n· 成交复盘 → 同类报价锚点 + 复购排程\n· 赢输单原话自动回写进租户知识库（买家语言 / 标准异议）',
    13.5,WHITE,False,spacing=1.3)
txt(s,0.9,5.6,11.5,1.2,
    '案例实证：成交 ¥12.6 万年度系统单，复盘出「需求变更未设闸门 + 验收未锁死」两类根因 → 沉淀为变更收费/验收即锁 SOP → 客户满意转复购。\n知识随业务自动生长：人不必手工维护话术库，复盘本身就是知识更新。',
    13.5,GREEN,False,spacing=1.3)

# =====================================================================
# 11 · 平台亮点二：多行业零代码配置上线
# =====================================================================
s = newslide()
header(s,'05 · 平台亮点 ②','多行业零代码配置上线：一个行业 = 一份配置画像','实体 · 阶段 · 审批 · 关系 · 领域计算全部配置化，不改内核、不污染其他租户')
inds = [
    ('配置画像','粒子类型/属性/阶段/审批/关系，由租户级配置画像声明；新增行业不新增代码、不加字面上限。',CYAN),
    ('阶段与审批可配','阶段名称、数量、停留阈值、必做动作按行业自定义；审批域不再写死，销售阈值同样配置化。',AMBER),
    ('领域计算引擎','价格链路、毛利红线、分级规则由公式引擎承载，行业专属算法不进内核分支。',PURPLE),
    ('租户隔离','粒子、配置、Knowledge 全部按租户隔离；开通即播种行业主数据，停用只改状态不删数据。',GREEN),
]
positions = [(0.9,2.45),(6.5,2.45),(0.9,4.25),(6.5,4.25)]
for (t,d,c),(l,ty) in zip(inds, positions):
    card(s,l,ty,5.4,1.65,fill=BG2,line=c)
    txt(s,l+0.25,ty+0.2,5.0,0.5,t,17,c,True)
    txt(s,l+0.25,ty+0.85,5.0,0.7,d,12.5,GRAY,False,spacing=1.25)
txt(s,0.9,6.25,11.5,0.6,'行业差异留在配置里，不留在代码里 —— 同一套内核，可同时承载制造、家电、新能源、培训服务等不同业态。',14,CYAN,True)

# =====================================================================
# 12 · 全景：八大能力合成飞轮
# =====================================================================
s = newslide()
header(s,'06 · 全景','八大能力不是八个功能，而是一个能力飞轮','方法论 → 行为 → 流程 → 记忆 → 校验 → 决策 → 复盘 → 知识回写')
fly = [
    ('业务 · 方法论','15 个 method-* SKILL：判断有方法',BIZ),
    ('业务 · 行为习惯','21 条行为标准：动作有标尺',AMBER),
    ('业务 · LTC 内嵌','线索到回款一条轨道：流程不脱节',GREEN),
    ('技术 · 记忆系统','上下文图谱 + Know-How：客户不丢失',TECH),
    ('平台 · K-M-D 引擎','九尺子校验：判断可解释、可复现',CYAN),
    ('技术 · 决策留痕','第 0 闸 + 事件驱动：判断可溯源',BLUE),
    ('技术 · 每日复盘','retro 引擎：经验每天被沉淀',GREEN),
    ('平台 · 多行业配置','配置画像：行业零代码上线',AMBER),
]
positions = [(0.6,2.4),(3.8,2.4),(7.0,2.4),(10.2,2.4),
             (0.6,4.05),(3.8,4.05),(7.0,4.05),(10.2,4.05)]
for (t,d,c),(l,ty) in zip(fly, positions):
    card(s,l,ty,3.0,1.45,fill=BG2,line=c)
    txt(s,l+0.2,ty+0.16,2.65,0.45,t,14,c,True)
    txt(s,l+0.2,ty+0.7,2.65,0.65,d,11.5,GRAY,False,spacing=1.15)
txt(s,0.6,5.75,12.2,1.1,
    '越用越聪明：捕获 Capture → 拼装 Assemble → 校验 Verify → 决策 Decide → 回流 Feedback → 复盘 Review → 知识回写。\n每一圈，先例更丰富、判断更准、人更省心。',
    14.5,WHITE,False,spacing=1.35)

# =====================================================================
# 13 · 价值与落地：真实业务印证
# =====================================================================
s = newslide()
header(s,'06 · 价值与落地','用真实业务印证能力落地','报价 / 需求补全 / 链预警 / 决策可解释 / 角色自适应 / 复盘复购，全部来自真实业务')
proofs = [
    ('系统方案报价','一句话报出 A/B 两方案，成本结构、毛利底线全透视',GREEN),
    ('需求自动补全','8 项必填缺失自动追问，一个来回齐',BIZ),
    ('交付风险预警','链断裂/回款逾期先于提问主动推送',RED),
    ('决策可解释','九尺子逐项打分，不达标项归因后复核，不放行',TECH),
    ('角色自适应','同一句话，销售/经理/高管/财务各见所需',AMBER),
    ('订单复盘复购','¥12.6 万复盘成 SOP，复购自动排程',GREEN),
]
positions = [(0.9,2.4),(6.5,2.4),(0.9,4.15),(6.5,4.15),(0.9,5.9),(6.5,5.9)]
for (t,d,c),(l,ty) in zip(proofs, positions):
    card(s,l,ty,5.4,1.55,fill=CARD,line=c)
    txt(s,l+0.25,ty+0.16,5.0,0.45,t,15,c,True)
    txt(s,l+0.25,ty+0.7,5.0,0.75,d,12.5,GRAY,False,spacing=1.15)

# =====================================================================
# 14 · 结语
# =====================================================================
s = newslide()
header(s,'结语 · WHAT\'S NEXT','让「方法、习惯、流程、记忆、校验、复盘、知识」都成为系统能力','建议下一步：用真实数据说话')
txt(s,0.9,2.5,11.5,1.2,
    '平台不是工具，而是会自检、会校验、会进化的销售中枢。\n先跑通「一般商机全自动闭环」：自动接诊、九尺子校验、标准报价、事件驱动跟进，让销售立刻减负。',
    16.5,WHITE,False,spacing=1.35)
pill(s,0.9,4.15,3.5,0.6,'第一步 · 一键召唤专家包',CYAN,BG2,13)
pill(s,4.7,4.15,3.8,0.6,'第二步 · 方法论+行为标准铺开',AMBER,BG2,13)
pill(s,8.6,4.15,3.8,0.6,'第三步 · 每日复盘自动运转',GREEN,BG2,13)
txt(s,0.9,5.45,11.5,1.2,
    '企业AI销售决策平台 —— 让每一次报价有方法、每一次拜访有标准、\n每一个决策有依据、每一天的经验都被沉淀、换行业只换一份配置。',
    17,AMBER,True,spacing=1.3)

# ---------- 保存 ----------
out = 'D:/system/CRM-ai-native/doc/企业AI销售决策平台介绍_2026_v2.pptx'
try:
    prs.save(out)
    print('SAVED_OK', out, 'slides=', len(prs.slides._sldIdLst))
except PermissionError:
    out2 = out.replace('.pptx', '_new.pptx')
    prs.save(out2)
    print('LOCKED_SAVED_AS', out2, 'slides=', len(prs.slides._sldIdLst))
