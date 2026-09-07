# -*- coding: utf-8 -*-
"""直接修改 Template.docx：裁剪两段超限文案 + 追加线上报名表补充字段部分（python-docx 直写）"""
import sys, io, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Pt, RGBColor

PATH = r"C:\Users\wangchuan08\Desktop\Template.docx"
BAK  = r"D:\system\CRM-ai-native\doc\event\Template.backup.docx"

shutil.copy2(PATH, BAK)
print("backup ->", BAK)

OLD_INNOV = "①15 个销售方法论 SKILL 化，每次判断自动装配，不再拍脑袋；②K-M-D 认知决策引擎，九尺子校验中 8 项走确定性规则、可复现；③写时向量化客户记忆，零填表秒级命中；④写操作必经决策第 0 闸，全程留痕可审计；⑤多行业零代码配置上线，换行业不改代码。"
NEW_INNOV = "①15 个方法论 SKILL 化，判断自动装配；②K-M-D 引擎 8 项校验走确定性规则、可复现；③写时向量化记忆，零填表秒级命中；④写操作必经决策第 0 闸，全程留痕；⑤多行业零代码配置上线。"

OLD_EFF = "平台已在真实 B2B 销售业务运行，支撑多租户 SaaS 全流程：试点数据显示智能接诊平均响应 1.2 分钟、报价测算 42 秒、线索零漏接；商机阶段停留超阈值自动预警，超权限报价强制审批闭环，客户全景 10 秒可查，行为达标率实时可见，销售管理与合规风险显著下降。"
NEW_EFF = "平台已在真实 B2B 销售业务运行：智能接诊平均响应 1.2 分钟、报价测算 42 秒、线索零漏接；超权限报价强制审批闭环，客户全景 10 秒可查，行为达标率实时可见，管理与合规风险显著下降。"

doc = Document(PATH)

replaced = {"inov": 0, "eff": 0}
def replace_in_paragraph(p, old, new):
    if old in p.text:
        # 保留首个 run 的格式，重写整段文本
        full = p.text.replace(old, new)
        for r in list(p.runs):
            r.text = ""
        if p.runs:
            p.runs[0].text = full
        else:
            p.add_run(full)
        return True
    return False

def iter_all_paragraphs(d):
    for p in d.paragraphs:
        yield p
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
                for nt in cell.tables:
                    for nrow in nt.rows:
                        for ncell in nrow.cells:
                            for np in ncell.paragraphs:
                                yield np

for p in iter_all_paragraphs(doc):
    if replace_in_paragraph(p, OLD_INNOV, NEW_INNOV):
        replaced["inov"] += 1
    if replace_in_paragraph(p, OLD_EFF, NEW_EFF):
        replaced["eff"] += 1

print("replaced 主要创新 x%d, 应用成效 x%d" % (replaced["inov"], replaced["eff"]))

# 追加「七、线上报名表补充字段」
PARAS = [
    ("项目所用算力", "项目算力策略为“确定性规则前置＋LLM 按需调用”——九尺子校验 8 项走代码判定、语义检索走 pgvector 向量索引，仅方法论推理与复盘归因调用大模型（接入主流 OpenAI 兼容 API，模型可插拔）。开发与生产均基于 x86_64 通用服务器与云端 CPU/GPU 弹性算力，无专用加速卡依赖；单租户 4 核 8G 云主机即可承载百人级销售团队日常运行，随租户增长水平扩展。"),
    ("建设成本（万元）", "约 300（按累计研发投入口径：60290 行自研代码＋团队人力＋云资源。注：此为建议口径，提交前请按真实数字核定修改）。"),
    ("经济效益（万元）", "已实现收益据实填写；若暂未产生收益填 0（提交前请核定）。"),
    ("是否使用绿色算力", "部分使用——生产环境部署于云服务商数据中心，绿电占比以云厂商年度绿电披露为准；平台通过“确定性规则前置”架构显著减少大模型调用量，单位决策能耗低于同类全大模型方案，符合绿色算力导向。"),
    ("项目代表图片建议", "图片一（项目宣传图）用封面宣传图；图片二（系统架构图）用技术架构图；其他可补产品功能总览、商业模式图、市场定位图及产品界面截图（均见 doc/event/bp/ 目录）。"),
]

h = doc.add_heading("", level=1)
h.add_run("七、线上报名表补充字段（步骤 2 项目概要之外）")
for title, body in PARAS:
    p = doc.add_paragraph()
    r = p.add_run(title + "：")
    r.bold = True
    p.add_run(body)

doc.save(PATH)
print("saved ->", PATH)
